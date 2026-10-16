"""
Master Production Formatter for Digit_Recognition_AIML-B.docx
Transforms the document to 100% compliance with:
1. Institutional Guidelines:
   - Paper size: A4 throughout (210 mm x 297 mm)
   - Margins: Top=0.9", Bottom=0.9", Left=1.0", Right=1.0", Header=0.45", Footer=0.45"
   - Institution Header on every page: 4-line header in Times New Roman
   - Running Footer: Left text 'Micro Project Report | Department of CSE (AI & ML)' + Right dynamic PAGE field
   - Page Numbering: Lowercase Roman (i, ii, iii..) for Preliminaries, Decimal (1, 2, 3..) for Main
   - Fonts: Times New Roman throughout; Code in Courier New 9.5-10 pt
   - Font sizes: Chapter titles 16 pt bold; section headings 13 pt bold; sub-headings 12 pt bold; body text 12 pt; tables 10-11 pt
   - Line spacing: 1.5 for body text; Single (1.0) within tables
   - Figures & Tables: Table captions ABOVE (11 pt bold), Figure captions BELOW (11 pt bold)
   - Citations: IEEE style, numbered in square brackets [1]..[10]
   - Results: Empirical numbers to 3-4 decimal places from test set
2. Abhinav Orientation Document Layout:
   - Order: Title Sheet (i) -> Certificate (ii) -> Declaration & Acknowledgement (iii) ->
            Abstract (iv) -> Outcome Mapping (v-vii) -> Table of Contents (1) -> Chapters 1-9 (2-34) ->
            References (36) -> Appendices A, B, C (37-39)
   - Dignitaries: Dr. A. Swathi (Coordinator, Guide, HOD), Dr. C. Udaya Kiran (Principal), Mr. A. Hriday Reddy (Vice-Chairman)
   - Course Details: Course Code: A1612, Machine Learning Laboratory, B.Tech III Year I Semester, Section: B, Academic Year: 2026 – 2027, Regulation: R22
   - Team: M TEJAS YADAV (24VE1A6695), K Charan Teja (24VE1A6692), Y Akshay Raj (24VE1A66C8)
"""

import os
import re
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    """Set cell internal padding in twips/dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(child)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="E8E8E8"):
    """Set cell background fill color."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in tcPr.findall(qn('w:shd')):
        tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def apply_run_font(run, font_name="Times New Roman", size_pt=12.0, bold=None, italic=None, color=None):
    """Set exact font typography on run and OpenXML rFonts."""
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color

def apply_paragraph_style(p, font_name="Times New Roman", size_pt=12.0, bold=None, italic=None,
                          color=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5,
                          space_before=0, space_after=4, keep_with_next=False, page_break_before=False):
    """Format paragraph margins, alignment, line spacing, and typography."""
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = keep_with_next
    if page_break_before:
        p.paragraph_format.page_break_before = True

    for r in p.runs:
        apply_run_font(r, font_name=font_name, size_pt=size_pt, bold=bold, italic=italic, color=color)

def setup_header(header):
    """Set up the standard 4-line institutional header in Times New Roman."""
    for p in header.paragraphs:
        p._p.getparent().remove(p._p)
        
    lines = [
        ("Sreyas Institute of Engineering and Technology", 16.0, True, 0, 1),
        ("An Autonomous Institution", 12.0, True, 0, 1),
        ("Approved by AICTE, Affiliated to JNTUH", 10.0, False, 0, 1),
        ("Accredited by NAAC-A Grade, NBA (CSE, ECE & ME) & ISO 9001:2015 Certified", 10.0, False, 0, 2)
    ]
    for text, sz, bld, sb, sa in lines:
        p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(sb)
        p.paragraph_format.space_after = Pt(sa)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        apply_run_font(r, font_name="Times New Roman", size_pt=sz, bold=bld)

def setup_footer(footer, section, is_roman=False):
    """Set up the running footer with left text and right dynamic PAGE field."""
    for p in footer.paragraphs:
        p._p.getparent().remove(p._p)
        
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    # 451.3 pt is text width on A4 with 1-inch margins
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Pt(451.3), WD_TAB_ALIGNMENT.RIGHT)
    
    r_left = p.add_run("Micro Project Report | Department of CSE (AI & ML)")
    apply_run_font(r_left, font_name="Times New Roman", size_pt=9.0, bold=False)
    
    r_tab = p.add_run("\t")
    apply_run_font(r_tab, font_name="Times New Roman")
    
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p._p.append(fldSimple)
    
    sectPr = section._sectPr
    for child in sectPr.findall(qn('w:pgNumType')):
        sectPr.remove(child)
    fmt = "lowerRoman" if is_roman else "decimal"
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}" w:start="1"/>')
    sectPr.append(pgNumType)

def format_table(table, header_bg="E8E8E8", is_code=False):
    """Format table cells: 10-11 pt, single line spacing, Times New Roman (or Courier New for code)."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    font_name = "Courier New" if is_code else "Times New Roman"
    
    for r_idx, row in enumerate(table.rows):
        is_header = (r_idx == 0 and not is_code and len(table.rows) > 1)
        for cell in row.cells:
            set_cell_margins(cell, top=65, bottom=65, left=95, right=95)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if is_header:
                set_cell_shading(cell, header_bg)
            
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.0 # strictly single within tables
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                
                for r in p.runs:
                    if is_code:
                        apply_run_font(r, font_name="Courier New", size_pt=9.5, bold=False)
                    elif is_header:
                        apply_run_font(r, font_name="Times New Roman", size_pt=10.5, bold=True)
                    else:
                        apply_run_font(r, font_name="Times New Roman", size_pt=10.0)

def execute_master_remediation():
    src_path = "Digit_Recognition_AIML-B_Backup.docx"
    target_path = "Digit_Recognition_AIML-B.docx"
    print(f"Loading clean backup from {src_path}...")
    doc = docx.Document(src_path)

    # -----------------------------------------------------------------
    # STEP 1: Re-order TOC to follow Outcome Mapping (Matching Abhinav)
    # -----------------------------------------------------------------
    print("Re-ordering Table of Contents to follow Outcome Mapping (Abhinav layout)...")
    body = doc._element.body
    toc_p = None
    toc_tbl = None
    pso_tbl = None
    p_sect = None

    for child in body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p_obj = docx.text.paragraph.Paragraph(child, doc)
            if 'TABLE OF CONTENTS' in p_obj.text:
                toc_p = child
            if child.xpath('.//w:sectPr'):
                p_sect = child
        elif tag == 'tbl':
            tbl = docx.table.Table(child, doc)
            if tbl.rows and 'Chapter' in tbl.rows[0].cells[0].text:
                toc_tbl = child
            if tbl.rows and 'PSO' in tbl.rows[0].cells[0].text:
                pso_tbl = child

    if toc_p is not None and toc_tbl is not None and pso_tbl is not None:
        body.remove(toc_p)
        body.remove(toc_tbl)
        # Place TOC elements after pso_tbl
        pso_tbl.addnext(toc_tbl)
        pso_tbl.addnext(toc_p)
        print("  Table of Contents successfully positioned after Outcome Mapping!")

    # -----------------------------------------------------------------
    # STEP 2: Configure Sections, A4 Dimensions, Margins, Header/Footer
    # -----------------------------------------------------------------
    print("Setting up Sections: A4 dimensions, margins, running header and footer...")
    for idx, s in enumerate(doc.sections):
        s.page_width = Pt(595.3)   # 210 mm
        s.page_height = Pt(841.9)  # 297 mm
        s.top_margin = Pt(64.8)    # 0.9 inch
        s.bottom_margin = Pt(64.8) # 0.9 inch
        s.left_margin = Pt(72.0)   # 1.0 inch
        s.right_margin = Pt(72.0)  # 1.0 inch
        s.header_distance = Pt(32.4) # 0.45 inch
        s.footer_distance = Pt(32.4) # 0.45 inch
        s.different_first_page_header_footer = False
        
        setup_header(s.header)
        is_roman = (idx == 0)
        setup_footer(s.footer, s, is_roman=is_roman)

    # -----------------------------------------------------------------
    # STEP 3: Split Merged Headings & Glitched Paragraphs
    # -----------------------------------------------------------------
    print("Splitting merged headings and paragraphs...")
    split_targets = [
        ("1.5  Scope and Limitations", "The scope of this project encompasses"),
        ("1.6  Organisation of the Report", "This micro project report is structured"),
        ("2.2Research Gap Identified", "A critical review of existing academic"),
        ("2.2  Research Gap Identified", "A critical review of existing academic"),
        ("2.3  Proposed Approach", "DigitVision AI bridges these research gaps"),
        ("9.2  Limitations", "Current platform limitations include:"),
        ("9.3Future Scope", "Future enhancements will focus on"),
        ("9.3  Future Scope", "Future enhancements will focus on")
    ]
    
    for p in doc.paragraphs:
        txt = p.text
        for h_str, b_str in split_targets:
            if h_str in txt and b_str in txt:
                clean_h = h_str.replace("2.2Research", "2.2  Research").replace("9.3Future", "9.3  Future")
                parts = txt.split(b_str)
                p.text = clean_h
                p.style = "Heading 2"
                apply_paragraph_style(p, size_pt=13.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                                      line_spacing=1.5, space_before=12, space_after=4, keep_with_next=True)
                
                new_p = doc.add_paragraph()
                new_p.text = b_str + parts[1]
                p._p.addnext(new_p._p)
                apply_paragraph_style(new_p, size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                                      line_spacing=1.5, space_before=0, space_after=4)
                print(f"  Split: '{clean_h}'")
                break

    # -----------------------------------------------------------------
    # STEP 4: Front Matter Harmonization (Preliminaries)
    # -----------------------------------------------------------------
    print("Formatting Preliminaries (Cover, Certificate, Declaration, Acknowledgement, Abstract)...")
    
    # Title Sheet Paragraphs
    doc.paragraphs[0].text = "Department of CSE (Artificial Intelligence & Machine Learning)"
    apply_paragraph_style(doc.paragraphs[0], size_pt=14.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=8, space_after=8)

    doc.paragraphs[1].text = "MICRO PROJECT REPORT"
    apply_paragraph_style(doc.paragraphs[1], size_pt=20.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=4, space_after=4)

    doc.paragraphs[2].text = "on"
    apply_paragraph_style(doc.paragraphs[2], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0, space_before=2, space_after=4)

    doc.paragraphs[3].text = "DIGITVISION AI: EXPLAINABLE, CONFIDENCE-AWARE HANDWRITTEN DIGIT INTELLIGENCE PLATFORM"
    apply_paragraph_style(doc.paragraphs[3], size_pt=14.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.2, space_before=4, space_after=12)

    doc.paragraphs[4].text = "Submitted in partial fulfilment of the requirements of the course"
    apply_paragraph_style(doc.paragraphs[4], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=4, space_after=2)

    doc.paragraphs[5].text = "Course Code: A1612"
    apply_paragraph_style(doc.paragraphs[5], size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=1, space_after=1)

    doc.paragraphs[6].text = "Course Name: Machine Learning Laboratory\nB.Tech III Year I Semester   |   Section: B\nAcademic Year: 2026 – 2027   |   Regulation: R22"
    apply_paragraph_style(doc.paragraphs[6], size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=1, space_after=4)

    doc.paragraphs[7].text = "Submitted by"
    apply_paragraph_style(doc.paragraphs[7], size_pt=12.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=8, space_after=4)

    doc.paragraphs[8].text = ""
    apply_paragraph_style(doc.paragraphs[8], size_pt=2.0, line_spacing=1.0, space_before=0, space_after=0)

    doc.paragraphs[10].text = "Under the Guidance of\nDr. A. Swathi\nAssociate Professor, Department of CSE (AI & ML)\nMonth, Year of Submission: October, 2026"
    apply_paragraph_style(doc.paragraphs[10], size_pt=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.2, space_before=8, space_after=8)

    # Certificate
    doc.paragraphs[11].text = "CERTIFICATE"
    apply_paragraph_style(doc.paragraphs[11], size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                          line_spacing=1.15, space_before=16, space_after=12, page_break_before=True)

    doc.paragraphs[14].text = (
        "This is to certify that the micro project entitled \"DigitVision AI: Explainable, Confidence-Aware "
        "Handwritten Digit Intelligence Platform\" is a Bonafide work carried out by M TEJAS YADAV (Roll No. 24VE1A6695), "
        "K Charan Teja (Roll No. 24VE1A6692), and Y Akshay Raj (Roll No. 24VE1A66C8) of B.Tech III Year I Semester, "
        "Department of CSE (Artificial Intelligence & Machine Learning), Sreyas Institute of Engineering and Technology, "
        "in partial fulfilment of the requirements of the course Machine Learning Laboratory (Course Code: A1612) "
        "during the academic year 2026–2027. The work has been carried out under my supervision and has been evaluated "
        "using the departmental assessment rubrics."
    )
    apply_paragraph_style(doc.paragraphs[14], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_before=6, space_after=12)

    doc.paragraphs[16].text = "Marks awarded: ______ / 10\nEvaluation date: 13-10-2026\n\nSignature\nCourse Coordinator\nName: Dr. A Swathi\nDate: 13-10-2026"
    apply_paragraph_style(doc.paragraphs[16], size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.2, space_before=10, space_after=10)

    # Declaration
    doc.paragraphs[18].text = "DECLARATION"
    apply_paragraph_style(doc.paragraphs[18], size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                          line_spacing=1.15, space_before=16, space_after=10, page_break_before=True)

    doc.paragraphs[19].text = (
        "We hereby declare that the micro project report entitled \"DigitVision AI: Explainable, Confidence-Aware "
        "Handwritten Digit Intelligence Platform\" submitted to the Department of CSE (AI & ML), Sreyas Institute "
        "of Engineering and Technology, is a record of original work done by us under the guidance of Dr. A. Swathi, "
        "Associate Professor, Department of CSE (AI & ML). The content of this report has not been submitted elsewhere "
        "for the award of any credit. All sources of information have been duly acknowledged and cited, and no part "
        "of this work is plagiarized."
    )
    apply_paragraph_style(doc.paragraphs[19], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_before=6, space_after=10)

    doc.paragraphs[21].text = "Place: Hyderabad                                                   Date: 13-10-2026"
    apply_paragraph_style(doc.paragraphs[21], size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15, space_before=6, space_after=12)

    # Acknowledgement
    doc.paragraphs[22].text = "ACKNOWLEDGEMENT"
    apply_paragraph_style(doc.paragraphs[22], size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.15, space_before=12, space_after=8)

    doc.paragraphs[23].text = (
        "We would like to express our sincere gratitude to our project guide, Dr. A. Swathi, Associate Professor, "
        "Department of Computer Science and Engineering (Artificial Intelligence & Machine Learning), for her invaluable "
        "guidance, constructive suggestions, and continuous encouragement throughout the successful completion of this machine learning project.\n\n"
        "We extend our heartfelt thanks to Dr. A. Swathi, Head of the Department of Computer Science and Engineering "
        "(Artificial Intelligence & Machine Learning), for her support and for providing a conducive academic environment to carry out this project.\n\n"
        "We express our sincere gratitude to Dr. C. Udaya Kiran, Principal, and Mr. A. Hriday Reddy, Vice-Chairman, "
        "for providing the necessary infrastructure, laboratory facilities, and institutional support required for the successful execution of this project.\n\n"
        "We would also like to thank all the faculty members, technical staff, and laboratory personnel of Sreyas Institute "
        "of Engineering and Technology for their valuable assistance and cooperation throughout the project. Finally, we are "
        "grateful to our peers and everyone who directly or indirectly contributed to the successful completion of this project."
    )
    apply_paragraph_style(doc.paragraphs[23], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_before=4, space_after=10)

    # Abstract
    doc.paragraphs[29].text = "ABSTRACT"
    apply_paragraph_style(doc.paragraphs[29], size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                          line_spacing=1.15, space_before=16, space_after=12, page_break_before=True)

    apply_paragraph_style(doc.paragraphs[30], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_before=6, space_after=10)

    doc.paragraphs[34].text = "Keywords: Handwritten Digit Recognition, Convolutional Neural Networks, Explainable AI (Grad-CAM), Temperature Scaling, Model Calibration, Edge Robustness."
    apply_paragraph_style(doc.paragraphs[34], size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.5, space_before=8, space_after=12)
    for r in doc.paragraphs[34].runs:
        if "Keywords:" in r.text:
            r.font.bold = True

    # Outcome Mapping Heading
    for p in doc.paragraphs:
        if "OUTCOME MAPPING" in p.text:
            p.text = "OUTCOME MAPPING (CO — PO — PSO)"
            apply_paragraph_style(p, size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                                  line_spacing=1.15, space_before=16, space_after=12, page_break_before=True)

    # Table of Contents Heading
    for p in doc.paragraphs:
        if "TABLE OF CONTENTS" in p.text:
            p.text = "TABLE OF CONTENTS"
            apply_paragraph_style(p, size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                                  line_spacing=1.15, space_before=16, space_after=12, page_break_before=True)

    # Synchronize TOC Table data
    t4 = [t for t in doc.tables if len(t.rows) > 0 and "Chapter" in t.rows[0].cells[0].text][0]
    toc_data = [
        ("", "Title Sheet", "i"),
        ("", "Certificate", "ii"),
        ("", "Declaration and Acknowledgement", "iii"),
        ("", "Abstract", "iv"),
        ("", "Outcome Mapping (CO — PO — PSO)", "v"),
        ("", "Table of Contents", "1"),
        ("1", "Introduction", "2"),
        ("2", "Literature Survey", "5"),
        ("3", "Dataset Description", "8"),
        ("4", "Data Preprocessing and Feature Engineering", "13"),
        ("5", "Methodology and Model Design", "16"),
        ("6", "Model Training and Convergence Analysis", "20"),
        ("7", "Evaluation and Results", "26"),
        ("8", "Responsible AI: Societal, Environmental and Ethical Considerations", "33"),
        ("9", "Conclusion and Future Scope", "34"),
        ("", "References", "36"),
        ("A", "Appendix A — Assessment Rubrics", "37"),
        ("B", "Appendix B — Individual Assessment Sheet", "38"),
        ("C", "Appendix C — Student Self-Reflection", "39")
    ]
    while len(t4.rows) < len(toc_data) + 1:
        t4.add_row()
    while len(t4.rows) > len(toc_data) + 1:
        t4._tbl.remove(t4.rows[-1]._tr)

    for i, (ch, tit, pg) in enumerate(toc_data):
        row = t4.rows[i + 1]
        row.cells[0].text = ch
        row.cells[1].text = tit
        row.cells[2].text = pg

    # -----------------------------------------------------------------
    # STEP 5: Expand References into distinct IEEE paragraphs
    # -----------------------------------------------------------------
    print("Formatting IEEE references...")
    ref_heading_p = None
    ref_paras = []
    in_refs = False

    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "REFERENCES":
            ref_heading_p = p
            in_refs = True
            continue
        if in_refs and txt.startswith("APPENDIX"):
            in_refs = False
            continue
        if in_refs and txt:
            ref_paras.append(p)

    all_citations = []
    for rp in ref_paras:
        # Extract citations [X] ...
        items = re.findall(r'(\[\d+\].*?)(?=\[\d+\]|\Z)', rp.text, re.DOTALL)
        for item in items:
            clean_item = item.strip().replace('\n', ' ')
            clean_item = re.sub(r'\s+', ' ', clean_item)
            if clean_item:
                all_citations.append(clean_item)

    if len(all_citations) >= 10:
        print(f"  Found {len(all_citations)} citations. Converting to clean paragraphs...")
        # Clear old reference paragraphs except the first one
        for p in ref_paras[1:]:
            p._p.getparent().remove(p._p)
            
        current_p = ref_paras[0]
        current_p.text = all_citations[0]
        apply_paragraph_style(current_p, size_pt=11.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                              line_spacing=1.15, space_before=2, space_after=4)
        current_p.paragraph_format.left_indent = Inches(0.25)
        current_p.paragraph_format.first_line_indent = Inches(-0.25)

        for cit in all_citations[1:]:
            new_p = doc.add_paragraph()
            new_p.text = cit
            current_p._p.addnext(new_p._p)
            apply_paragraph_style(new_p, size_pt=11.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                                  line_spacing=1.15, space_before=2, space_after=4)
            new_p.paragraph_format.left_indent = Inches(0.25)
            new_p.paragraph_format.first_line_indent = Inches(-0.25)
            current_p = new_p

    # -----------------------------------------------------------------
    # STEP 6: Full Document Standardized Typography Pass
    # -----------------------------------------------------------------
    print("Executing full document typography pass...")
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            continue

        # Chapter Titles (16 pt bold, centered, page break before)
        if any(txt == f"CHAPTER {k}" for k in range(1, 10)) or txt in ["REFERENCES", "APPENDIX A — ASSESSMENT RUBRICS", "APPENDIX B — INDIVIDUAL ASSESSMENT SHEET", "APPENDIX C — STUDENT SELF-REFLECTION"]:
            p.style = "Heading 1"
            apply_paragraph_style(p, size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                                  line_spacing=1.15, space_before=18, space_after=4, keep_with_next=True, page_break_before=True)
            continue

        # Chapter Subtitle (e.g. INTRODUCTION) right after CHAPTER X
        if idx > 0 and doc.paragraphs[idx - 1].text.strip().startswith("CHAPTER "):
            apply_paragraph_style(p, size_pt=16.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                                  line_spacing=1.15, space_before=2, space_after=14, keep_with_next=True)
            continue

        # Section Headings (1.1, 1.2... 13 pt bold, left aligned)
        if re.match(r'^\d+\.\d+\s+[A-Z]', txt) and not re.match(r'^\d+\.\d+\.\d+', txt):
            p.style = "Heading 2"
            apply_paragraph_style(p, size_pt=13.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                                  line_spacing=1.5, space_before=12, space_after=4, keep_with_next=True)
            continue

        # Sub-section Headings (3.5.1... 12 pt bold, left aligned)
        if re.match(r'^\d+\.\d+\.\d+\s+[A-Z]', txt):
            p.style = "Heading 3"
            apply_paragraph_style(p, size_pt=12.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                                  line_spacing=1.5, space_before=8, space_after=3, keep_with_next=True)
            continue

        # Table Captions (11 pt bold, above table)
        if txt.lower().startswith('table ') and len(txt) < 120 and ':' in txt:
            apply_paragraph_style(p, size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                                  line_spacing=1.2, space_before=8, space_after=3, keep_with_next=True)
            continue

        # Figure Captions (11 pt bold, below figure)
        if (txt.lower().startswith('fig') or 'fig.' in txt.lower()) and len(txt) < 140:
            apply_paragraph_style(p, size_pt=11.0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                                  line_spacing=1.2, space_before=4, space_after=8)
            continue

        # Reference Citation Paragraphs
        if re.match(r'^\[\d+\]', txt):
            apply_paragraph_style(p, size_pt=11.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                                  line_spacing=1.15, space_before=2, space_after=4)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            continue

        # Standard Body Paragraphs (12 pt, 1.5 line spacing, justified)
        if idx > 34:
            apply_paragraph_style(p, size_pt=12.0, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                                  line_spacing=1.5, space_before=0, space_after=4)

    # -----------------------------------------------------------------
    # STEP 7: Format all 52 Tables (10-11 pt, single line spacing)
    # -----------------------------------------------------------------
    print("Formatting all tables (10-11 pt, single line spacing, Times New Roman / Courier New)...")
    for t_idx, table in enumerate(doc.tables):
        is_code = (t_idx == 35) # Key Code Segments
        format_table(table, header_bg="E8E8E8", is_code=is_code)

    # Save output
    doc.save(target_path)
    print(f"Master standardized document successfully saved to: {target_path}")

if __name__ == "__main__":
    execute_master_remediation()
